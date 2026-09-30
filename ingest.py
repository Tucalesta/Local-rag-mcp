import os
import stat
# Disabilitiamo l'import security check di NLTK nel caso in cui la venv risieda nella cartella del progetto
os.environ["NLTK_DISABLE_IMPORT_SECURITY"] = "1"

import sys
import json
import math
import re
import time
import shutil
import logging
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta

from llama_index.core import Settings, StorageContext, VectorStoreIndex
from llama_index.core.schema import Document
from llama_index.core.node_parser import SentenceSplitter
from llama_index.embeddings.ollama import OllamaEmbedding
from llama_index.llms.ollama import Ollama
from llama_index.vector_stores.chroma import ChromaVectorStore
import chromadb

# Didascalie immagini (prompt versionati, chiamata a Ollama, cache su disco):
# modulo condiviso con bench_captions.py
from captioning import (
    CaptionCache, DEFAULT_MAX_SIDE, DEFAULT_PROMPT_VERSION, caption_image, get_prompt,
    path_is_under, prune_cache_file,
)

# Configurazione, embedding e database delle localita' condivisi con il server MCP
from rag_common import (
    collection_name, embed_model_name, find_nearest_location, haversine_km,
    load_gpx_database, make_embed_model, open_collection,
)

# Configuriamo Ollama
Settings.embed_model = OllamaEmbedding(model_name="nomic-embed-text")
Settings.llm = Ollama(model="llama3", request_timeout=120.0)

IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".webp", ".bmp")

# Modello vision di default (sovrascrivibile da config.json con la chiave
# "vision_model", senza dover toccare lo script). Gira su Ollama, lo stesso
# runtime gia' usato per l'LLM e per nomic-embed-text.
OLLAMA_VISION_MODEL_DEFAULT = "llama3.2-vision"

# Nome del file che tiene traccia di cosa e' gia' stato indicizzato (path,
# data di modifica, dimensione), per permettere l'indicizzazione
# incrementale: ai run successivi si processano solo i file nuovi o
# cambiati, invece di ricostruire tutto da zero.
MANIFEST_FILENAME = "ingest_manifest.json"

# Cartella con i GPX "a punti" (non tracce) usati come database personale di
# localita': un <wpt> per comune/paese, il nome del file GPX e' la regione.
# Serve per il reverse geocoding delle foto, offline e senza limiti di
# richieste (a differenza dei servizi online come Nominatim).
GPX_DIR_NAME = "res"

# Versione del formato di testo e metadati scritti nell'indice. Va incrementata
# quando cambia cio' che viene indicizzato: i file gia' indicizzati con un'altra
# versione vengono reindicizzati (le didascalie restano in cache: costa solo
# l'embedding).
INDEX_TEXT_VERSION = 2


def load_manifest(manifest_path, logger):
    if not os.path.exists(manifest_path):
        return {}
    try:
        with open(manifest_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        logger.warning(f"Impossibile leggere il manifest precedente ({e}). Si riparte da zero.")
        return {}


def save_manifest(manifest_path, manifest, logger):
    try:
        manifest_dir = os.path.dirname(manifest_path)
        if manifest_dir and not os.path.exists(manifest_dir):
            os.makedirs(manifest_dir, exist_ok=True)
        with open(manifest_path, "w", encoding="utf-8") as f:
            json.dump(manifest, f, ensure_ascii=False, indent=2)
    except Exception as e:
        logger.error(f"Impossibile salvare il manifest di indicizzazione: {e}")


def get_file_signature(file_path):
    """Firma leggera di un file (data di modifica + dimensione) per capire se
    e' cambiato dall'ultima indicizzazione, senza dover rileggere/hashare
    l'intero contenuto (costoso su archivi grandi)."""
    st = os.stat(file_path)
    return {"mtime": st.st_mtime, "size": st.st_size}


def _convert_gps_coord(value, ref):
    """Converte una coordinata GPS EXIF (gradi, minuti, secondi + riferimento
    N/S/E/W) nel formato decimale standard (es. 45.464201)."""
    if value is None or ref is None:
        return None
    try:
        degrees, minutes, seconds = value
        decimal = float(degrees) + float(minutes) / 60 + float(seconds) / 3600
        if ref in ("S", "W"):
            decimal = -decimal
        return round(decimal, 6)
    except Exception:
        return None


def extract_image_exif(file_path):
    """
    Estrae i metadati EXIF piu' utili da una foto, se presenti: data dello
    scatto, coordinate GPS (convertite in gradi decimali) e modello della
    fotocamera. Restituisce un dizionario "piatto" (solo str/int/float),
    compatibile con i metadata di ChromaDB. Le icone/grafiche generate non
    hanno quasi mai EXIF, quindi qui torna quasi sempre vuoto per loro -
    questa funzione si attiva di fatto solo sulle foto vere.
    """
    exif_data = {}
    try:
        from PIL import Image
        from PIL.ExifTags import TAGS, GPSTAGS

        with Image.open(file_path) as img:
            raw_exif = img.getexif()
            if not raw_exif:
                return exif_data

            for tag_id, value in raw_exif.items():
                tag = TAGS.get(tag_id, tag_id)
                if tag in ("DateTime", "DateTimeOriginal") and "date_taken" not in exif_data:
                    exif_data["date_taken"] = str(value).replace("\x00", "").strip()
                elif tag == "Make":
                    exif_data["camera_make"] = str(value).replace("\x00", "").strip()
                elif tag == "Model":
                    exif_data["camera_model"] = str(value).replace("\x00", "").strip()

            # DateTimeOriginal (momento dello scatto) sta nel sotto-blocco Exif
            # (0x8769), non in quello principale: e' piu' affidabile di DateTime,
            # che i programmi di fotoritocco aggiornano alla data di salvataggio.
            try:
                original = raw_exif.get_ifd(0x8769).get(36867)
                if original:
                    exif_data["date_taken"] = str(original).replace("\x00", "").strip()
            except Exception:
                pass

            # Le informazioni GPS sono in un IFD (sotto-blocco) separato,
            # identificato dal tag standard EXIF 0x8825 (GPSInfo).
            try:
                gps_ifd = raw_exif.get_ifd(0x8825)
            except Exception:
                gps_ifd = None

            if gps_ifd:
                gps_data = {GPSTAGS.get(k, k): v for k, v in gps_ifd.items()}
                lat = _convert_gps_coord(gps_data.get("GPSLatitude"), gps_data.get("GPSLatitudeRef"))
                lon = _convert_gps_coord(gps_data.get("GPSLongitude"), gps_data.get("GPSLongitudeRef"))
                if lat is not None and lon is not None:
                    exif_data["gps_lat"] = lat
                    exif_data["gps_lon"] = lon
    except Exception:
        # Formato senza EXIF, file corrotto, ecc.: nessun dato extra, non e' un errore.
        pass
    return exif_data


def is_likely_icon_or_graphic(file_path, min_side=128, max_distinct_colors=4096):
    """
    Euristica per riconoscere icone/pittogrammi/stencil (es. librerie di
    simboli per diagrammi tipo Dia/Visio) e distinguerli da foto reali,
    SENZA fare l'embedding costoso. Due segnali, sufficiente uno solo:

    1. Risoluzione molto piccola (icone tipiche: 16/32/48/64/96/128 px).
    2. Pochi colori distinti: una foto reale ha quasi sempre migliaia di
       sfumature diverse; un'icona piatta ne usa in genere poche decine.
       Questo secondo segnale funziona bene anche su icone grandi.
    """
    try:
        from PIL import Image
        with Image.open(file_path) as img:
            width, height = img.size
            if max(width, height) < min_side:
                return True

            # getcolors restituisce None se il numero di colori distinti
            # supera maxcolors: usiamo questo per un check economico senza
            # dover contare tutti i colori esplicitamente.
            rgb = img.convert("RGB")
            colors = rgb.getcolors(maxcolors=max_distinct_colors)
            if colors is not None:
                return True
    except Exception:
        # In caso di dubbio (immagine corrotta, formato strano) non blocchiamo:
        # la si prova comunque a embeddare, fallira' semmai piu' avanti.
        return False
    return False


def is_system_or_hidden_name(path):
    """
    Versione leggera del controllo, pensata per le cartelle RADICE indicate
    esplicitamente dall'utente in config.json (es. 'E:/'). Controlla solo il
    nome (Cestino, System Volume Information, ecc.) senza guardare gli
    attributi Windows SYSTEM/HIDDEN: la radice di un disco ha spesso
    l'attributo SYSTEM impostato da Windows per via del desktop.ini legato
    all'icona del disco, pur essendo una cartella scelta di proposito
    dall'utente - non va quindi esclusa solo per questo.
    """
    basename = os.path.basename(path.rstrip("/\\")).lower()
    system_names = {
        "$recycle.bin",
        "system volume information",
        "$winreagent",
        "config.msi",
        "found.000",
        ".trash",
        ".trashes"
    }
    return basename in system_names or basename.startswith("$")


def is_system_or_hidden(path):
    """
    Verifica se una cartella o un file è di sistema o nascosto (es. il Cestino $RECYCLE.BIN su Windows,
    System Volume Information, o elementi contrassegnati con attributi SYSTEM/HIDDEN).
    Usata per le SOTTOCARTELLE/file incontrati durante la scansione (non per
    la cartella radice configurata dall'utente: per quella si usa la
    versione leggera is_system_or_hidden_name, vedi sopra).
    """
    basename = os.path.basename(path).lower()

    # 1. Nomi noti di cartelle di sistema / Cestino
    system_names = {
        "$recycle.bin", 
        "system volume information", 
        "$winreagent", 
        "config.msi", 
        "found.000",
        ".trash",
        ".trashes"
    }

    if basename in system_names or basename.startswith("$"):
        return True

    # 2. Controllo degli attributi nativi Windows (SYSTEM e HIDDEN)
    try:
        st = os.stat(path)
        if hasattr(st, "st_file_attributes"):
            attrs = st.st_file_attributes
            if bool(attrs & stat.FILE_ATTRIBUTE_SYSTEM) or bool(attrs & stat.FILE_ATTRIBUTE_HIDDEN):
                return True
    except Exception:
        # Se non si hanno i permessi di lettura per lo stat, è molto probabilmente una cartella protetta di sistema
        return True

    return False

def parse_exif_date(value):
    """'2014:08:12 10:31:04' -> (2014, timestamp). None se assente o non valida
    (le fotocamere senza orologio impostato scrivono '0000:00:00 00:00:00')."""
    if not value:
        return None
    text = str(value).strip()
    for fmt, size in (("%Y:%m:%d %H:%M:%S", 19), ("%Y:%m:%d", 10), ("%Y-%m-%d %H:%M:%S", 19), ("%Y-%m-%d", 10)):
        try:
            dt = datetime.strptime(text[:size], fmt)
        except ValueError:
            continue
        if 1970 <= dt.year <= datetime.now().year + 1:
            try:
                return dt.year, int(dt.timestamp())
            except (OverflowError, OSError, ValueError):
                return None
        return None
    return None


def folder_label(folder, root_path):
    """Nomi delle cartelle tra la radice scansionata e la foto, es. '2014 / Dolomiti'.
    Utili per la ricerca (spesso dicono evento o luogo); niente lettera del disco,
    e ripetizioni consecutive ('2014 / 2014') ridotte a una."""
    try:
        rel = os.path.relpath(folder, root_path)
    except ValueError:
        return ""
    parts = [p for p in re.split(r"[\\/]+", rel) if p and p != "."]
    out = []
    for p in parts:
        if not out or out[-1].lower() != p.lower():
            out.append(p)
    return " / ".join(out)


def year_from_folder(label):
    """Anno (1900-2099) dal primo nome di cartella che ne contiene uno: serve per
    le foto senza data EXIF, tipiche degli scanner e delle vecchie fotocamere."""
    for part in label.split(" / ") if label else []:
        m = re.search(r"(?<!\d)(19\d{2}|20\d{2})(?!\d)", part)
        if m:
            return int(m.group(1))
    return None


def camera_label(make, model):
    make, model = (make or "").strip(), (model or "").strip()
    if make and model.lower().startswith(make.lower()):
        return model
    return f"{make} {model}".strip()


def count_candidate_images(config):
    """Conta le immagini che il ciclo di scansione esaminerebbe, guardando solo i
    nomi dei file (senza aprirle e senza interrogare gli attributi di ogni file,
    lento sui dischi di rete: e' solo una stima per l'avanzamento). Serve a
    mostrare l'avanzamento e il tempo residuo in modalita' solo-didascalie."""
    total = 0
    for entry in config.get("directories", []):
        target = entry.get("path")
        if not entry.get("index_images", True):
            continue
        if not target or not os.path.exists(target) or is_system_or_hidden_name(target):
            continue
        recursive = entry.get("recursive", False)
        for root, dirs, files in os.walk(target):
            dirs[:] = [d for d in dirs if not is_system_or_hidden(os.path.join(root, d))]
            total += sum(
                1 for f in files
                if f.lower().endswith(IMAGE_EXTENSIONS) and not is_system_or_hidden_name(f)
            )
            if not recursive:
                break
    return total


def setup_logging(logs_dir="./logs", retention_days=7):
    """
    Inizializza la registrazione dei log sia su console che su file nella cartella logs/.
    I file di log contengono la data nel nome (es. ingest_2026-08-09.log)
    e i file più vecchi di 7 giorni vengono cancellati automaticamente.
    """
    os.makedirs(logs_dir, exist_ok=True)

    # 1. Pulizia dei file di log più vecchi di 7 giorni
    now = datetime.now()
    cutoff_time = now - timedelta(days=retention_days)

    for filename in os.listdir(logs_dir):
        if filename.endswith(".log"):
            file_path = os.path.join(logs_dir, filename)
            try:
                file_mtime = datetime.fromtimestamp(os.path.getmtime(file_path))
                if file_mtime < cutoff_time:
                    os.remove(file_path)
                    print(f" -> Log datato rimosso (più vecchio di {retention_days} giorni): {filename}")
            except Exception as e:
                print(f" -> Errore durante la rimozione del vecchio log {filename}: {e}")

    # 2. Creazione del file di log per la data corrente
    date_str = now.strftime("%Y-%m-%d")
    log_filename = os.path.join(logs_dir, f"ingest_{date_str}.log")

    logger = logging.getLogger("ingest")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()  # Evita handlers duplicati se ri-eseguito

    formatter = logging.Formatter(
        fmt="[%(asctime)s] [%(levelname)s] %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S"
    )

    # Output su file
    file_handler = logging.FileHandler(log_filename, encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)

    # Output su console
    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(formatter)
    logger.addHandler(console_handler)

    return logger

def extract_text_from_file(file_path, ext, logger):
    """Estrae il testo in base all'estensione del file preservando la struttura a paragrafi."""
    text = ""
    try:
        # Se il file e' vuoto (0 byte), non c'e' contenuto da estrarre
        if os.path.getsize(file_path) == 0:
            logger.info(f"  -> File vuoto (0 byte) saltato: {os.path.basename(file_path)}")
            return ""

        if ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                logger.warning(f"Modulo 'pypdf' non installato. Impossibile leggere {file_path}. Installa con: pip install pypdf")
                return ""
            reader = PdfReader(file_path)
            pages_text = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(extracted)
            text = "\n\n".join(pages_text)
                    
        elif ext in [".txt", ".md"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
                
        elif ext == ".docx":
            try:
                from docx import Document as DocxDocument
            except ImportError:
                logger.warning(f"Modulo 'python-docx' non installato. Impossibile leggere {file_path}. Installa con: pip install python-docx")
                return ""
            try:
                doc = DocxDocument(file_path)
            except Exception as docx_err:
                logger.warning(f"Impossibile leggere il documento Word {file_path} (file non valido o corrotto): {docx_err}")
                return ""
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        paragraphs.append(" | ".join(row_text))
            text = "\n\n".join(paragraphs)
                        
        elif ext == ".xlsx":
            try:
                import openpyxl
            except ImportError:
                logger.warning(f"Modulo 'openpyxl' non installato. Impossibile leggere {file_path}. Installa con: pip install openpyxl")
                return ""
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheets_text = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                sheet_lines = [f"--- Foglio: {sheet} ---"]
                for row in ws.iter_rows(values_only=True):
                    row_values = [str(cell) for cell in row if cell is not None]
                    if row_values:
                        sheet_lines.append(" | ".join(row_values))
                sheets_text.append("\n".join(sheet_lines))
            text = "\n\n".join(sheets_text)
                        
    except Exception as e:
        logger.error(f"Errore di lettura per {file_path}: {e}")
        
    return text

def main():
    logger = setup_logging(logs_dir="./logs", retention_days=7)

    # BASE_DIR: cartella in cui si trova questo script. Serve per risolvere i
    # path relativi definiti in config.json indipendentemente dalla working
    # directory da cui viene lanciato ingest.py.
    base_dir = os.path.dirname(os.path.abspath(__file__))
    config_path = os.path.join(base_dir, "config.json")

    logger.info("=== Avvio processo di indicizzazione (ingest.py) ===")

    if not os.path.exists(config_path):
        logger.error(f"File di configurazione '{config_path}' non trovato.")
        return

    logger.info("1. Lettura della configurazione...")
    with open(config_path, "r", encoding="utf-8") as f:
        config = json.load(f)

    Settings.embed_model = make_embed_model(config)
    coll_name = collection_name(config)
    logger.info(f"  -> Modello di embedding: {embed_model_name(config)} | collezione ChromaDB: '{coll_name}'")

    # Modalita' solo-didascalie: genera e salva le didascalie nella cache, senza
    # toccare indice vettoriale e manifest e senza fare embedding. Utile quando
    # il modello vision e quello di embedding insieme non stanno in VRAM: prima
    # si descrivono tutte le foto, poi un secondo giro normale (senza
    # --caption-only) indicizza leggendo le didascalie dalla cache.
    caption_only = bool(config.get("caption_only", False)) or "--caption-only" in sys.argv[1:]
    if caption_only:
        logger.info("*** MODALITA' SOLO DIDASCALIE: nessuna modifica a indice e manifest ***")

    raw_db_path = config.get("db_path", "./chroma_db")
    if os.path.isabs(raw_db_path):
        db_path = raw_db_path
    else:
        db_path = os.path.normpath(os.path.join(base_dir, raw_db_path))
    logger.info(f"  -> Percorso database vettoriale: {db_path}")

    raw_manifest_path = config.get("manifest_path", config.get("manifest_file", MANIFEST_FILENAME))
    if os.path.isabs(raw_manifest_path):
        manifest_path = raw_manifest_path
    else:
        manifest_path = os.path.normpath(os.path.join(base_dir, raw_manifest_path))
    logger.info(f"  -> Percorso manifest di indicizzazione: {manifest_path}")

    vision_model = config.get("vision_model", OLLAMA_VISION_MODEL_DEFAULT)
    logger.info(f"  -> Modello vision per le didascalie immagini: {vision_model}")

    prompt_version = config.get("caption_prompt_version", DEFAULT_PROMPT_VERSION)
    try:
        get_prompt(prompt_version)
    except ValueError as e:
        logger.error(str(e))
        return
    caption_options = config.get("caption_options")
    caption_max_side = config.get("caption_max_side", DEFAULT_MAX_SIDE)
    caption_timeout = config.get("caption_timeout", 180)
    logger.info(f"  -> Prompt didascalie: versione '{prompt_version}', lato max immagine {caption_max_side} px")

    # Cache delle didascalie: se relativa, e' risolta rispetto alla cartella del
    # manifest (di norma la stessa del db). Sopravvive a force_reindex.
    raw_cache_path = config.get("caption_cache_path", "captions.jsonl")
    if os.path.isabs(raw_cache_path):
        caption_cache_path = raw_cache_path
    else:
        caption_cache_path = os.path.normpath(os.path.join(os.path.dirname(manifest_path), raw_cache_path))

    # Manutenzione della cache (ingest.py --prune-cache [--dry-run] [--only-current] [--force]):
    # non indicizza nulla, compatta captions.jsonl e termina.
    if "--prune-cache" in sys.argv[1:]:
        args = sys.argv[1:]
        roots = [e["path"] for e in config.get("directories", []) if e.get("path")]
        only_current = "--only-current" in args
        try:
            s = prune_cache_file(
                caption_cache_path, roots,
                only_model=vision_model if only_current else None,
                only_prompt=prompt_version if only_current else None,
                dry_run="--dry-run" in args, force="--force" in args,
            )
        except RuntimeError as e:
            logger.error(str(e))
            return
        logger.info(
            f"Cache didascalie {caption_cache_path}: {s['lines']} righe -> {s['kept']} conservate. "
            f"Rimosse: {s['duplicates']} duplicate, {s['orphans']} di file non piu' presenti, "
            f"{s['other_version']} di altro modello/prompt, {s['unreadable']} illeggibili."
        )
        if s["dry_run"]:
            logger.info("Simulazione (--dry-run): nessun file modificato.")
        elif s["backup"]:
            logger.info(f"Copia della cache precedente: {s['backup']}")
        else:
            logger.info("Niente da compattare.")
        return

    caption_cache = CaptionCache(caption_cache_path)
    logger.info(f"  -> Cache didascalie: {caption_cache_path} ({len(caption_cache)} voci)")

    raw_res_path = config.get("res_path", config.get("res_dir", config.get("gpx_dir", GPX_DIR_NAME)))
    if os.path.isabs(raw_res_path):
        res_dir = raw_res_path
    else:
        res_dir = os.path.normpath(os.path.join(base_dir, raw_res_path))
    logger.info(f"  -> Percorso database localita' (GPX/res): {res_dir}")

    gpx_max_km = config.get("gpx_max_distance_km", 15)
    gpx_points = load_gpx_database(res_dir, logger)
    geocoding_gaps = []

    force_reindex = config.get("force_reindex", False)
    old_manifest = {} if force_reindex else load_manifest(manifest_path, logger)
    new_manifest = {}
    unchanged_count = 0
    processed_count = 0
    supported_extensions = (".pdf", ".txt", ".md", ".docx", ".xlsx")

    # Checkpoint ibrido del manifest: si salva ogni 20 file processati OPPURE
    # ogni 5 minuti, quello che arriva prima. Scrivere il manifest ad ogni
    # singolo file sarebbe superfluo su archivi con molti file di testo
    # piccoli e veloci (I/O inutile); con questa soglia, nel caso peggiore si
    # perdono al massimo ~20 file o ~5 minuti di lavoro se il processo viene
    # interrotto, non l'intera run.
    CHECKPOINT_EVERY_N_FILES = 20
    CHECKPOINT_EVERY_SECONDS = 300
    _last_checkpoint_time = time.time()
    _files_since_checkpoint = 0

    def checkpoint():
        nonlocal _last_checkpoint_time, _files_since_checkpoint
        _files_since_checkpoint += 1
        elapsed = time.time() - _last_checkpoint_time
        if _files_since_checkpoint >= CHECKPOINT_EVERY_N_FILES or elapsed >= CHECKPOINT_EVERY_SECONDS:
            save_manifest(manifest_path, new_manifest, logger)
            _last_checkpoint_time = time.time()
            _files_since_checkpoint = 0

    # Apriamo il database vettoriale PRIMA di scansionare i file: ogni file
    # viene chunkato, embeddato, inserito e registrato nel manifest UNO ALLA
    # VOLTA (vedi index_single_document + i checkpoint dentro il ciclo).
    # Cosi', se l'esecuzione viene interrotta (Ctrl+C, chiusura del
    # terminale, crash) tutto cio' che e' stato completato fino a quel
    # momento resta acquisito: il prossimo ingest.py riparte esattamente da
    # dove ci si era fermati, invece di rifare da capo l'intera run.
    if caption_only:
        chroma_collection = index = text_splitter = None
    else:
        logger.info("2. Apertura del database vettoriale locale (ChromaDB)...")
        chroma_collection = open_collection(db_path, coll_name, create=True)
        vector_store = ChromaVectorStore(chroma_collection=chroma_collection)
        index = VectorStoreIndex.from_vector_store(vector_store=vector_store)
        text_splitter = SentenceSplitter(chunk_size=1000, chunk_overlap=200)
        # Manifest e collezione devono corrispondere: se il manifest elenca file
        # gia' indicizzati ma la collezione (nuova, per esempio dopo aver cambiato
        # embed_model) e' vuota, si riparte da zero invece di saltare tutto come
        # "invariato".
        if old_manifest and chroma_collection.count() == 0:
            logger.warning("  -> La collezione e' vuota ma il manifest elenca file gia' indicizzati: "
                           "ignoro il manifest e reindicizzo tutto.")
            old_manifest = {}

    def index_single_document(doc):
        """Chunka un Document, rimuove gli eventuali vecchi chunk dello stesso
        file (se gia' indicizzato in precedenza) e inserisce i nuovi chunk.
        Isolare l'operazione a un file per volta e' cio' che rende sicura
        un'interruzione a meta' processo."""
        path = doc.metadata.get("path")
        if path:
            try:
                chroma_collection.delete(where={"path": path})
            except Exception as e:
                logger.warning(f"  -> Impossibile ripulire i vecchi chunk di {path}: {e}")
        doc_nodes = text_splitter.get_nodes_from_documents([doc])
        if doc_nodes:
            index.insert_nodes(doc_nodes)
        return len(doc_nodes)

    # Statistiche della modalita' solo-didascalie.
    MAX_CONSECUTIVE_ERRORS = 5
    cap_stats = {"generated": 0, "cached": 0, "icons": 0, "empty": 0, "errors": 0,
                 "consecutive_errors": 0, "gen_seconds": 0.0}
    total_images = 0
    caption_only_started = time.time()

    def log_caption_only_summary():
        s = cap_stats
        logger.info(
            f"Solo didascalie: {s['generated']} generate in {(time.time() - caption_only_started) / 3600:.1f} h, "
            f"{s['cached']} gia' in cache, {s['icons']} icone/grafica saltate, {s['empty']} vuote, "
            f"{s['errors']} errori. Cache: {caption_cache_path} ({len(caption_cache)} voci)."
        )
        if s["generated"] or s["cached"]:
            logger.info("Prossimo passo: rilanciare ingest.py SENZA --caption-only per indicizzare "
                        "(le didascalie vengono lette dalla cache, il modello vision non serve).")

    if caption_only:
        logger.info("Conteggio delle foto da esaminare (solo nomi file)...")
        total_images = count_candidate_images(config)
        logger.info(f"  -> {total_images} foto candidate")

    logger.info("3. Scansione e indicizzazione incrementale dei file...")

    # Cartelle radice non raggiungibili (disco scollegato, rete spenta) o che
    # risultano senza nessun file da indicizzare (punto di montaggio vuoto): i
    # loro file gia' nell'indice NON vanno considerati rimossi.
    protected_roots = []
    healthy_roots = []

    for entry in config.get("directories", []):
        target_path = entry.get("path")
        recursive = entry.get("recursive", False)
        index_images = entry.get("index_images", True)
        root_files_seen = 0

        if not target_path or not os.path.exists(target_path):
            logger.warning(f"Percorso saltato (non esiste): {target_path}")
            if target_path:
                protected_roots.append(target_path)
            continue

        if is_system_or_hidden_name(target_path):
            logger.warning(f"Percorso saltato (cartella di sistema o nascosta): {target_path}")
            protected_roots.append(target_path)
            continue

        logger.info(f"Scansione cartella: {target_path} (Ricorsiva: {recursive})")

        if recursive:
            walk_target = os.walk(target_path, topdown=True)
        else:
            try:
                subdirs = []
                subfiles = []
                for item in os.listdir(target_path):
                    full_item = os.path.join(target_path, item)
                    if os.path.isdir(full_item):
                        subdirs.append(item)
                    else:
                        subfiles.append(item)
                walk_target = [(target_path, subdirs, subfiles)]
            except Exception as e:
                logger.error(f"Impossibile accedere alla cartella {target_path}: {e}")
                protected_roots.append(target_path)
                continue

        for root, dirs, files in walk_target:
            # Escludiamo in-place le cartelle di sistema e nascoste (es. $RECYCLE.BIN) prima che os.walk vi entri
            dirs[:] = [
                d for d in dirs
                if not is_system_or_hidden(os.path.join(root, d))
            ]

            for filename in files:
                file_path = os.path.join(root, filename)
                if not os.path.isfile(file_path) or is_system_or_hidden(file_path):
                    continue

                is_supported = filename.lower().endswith(supported_extensions)
                is_image = filename.lower().endswith(IMAGE_EXTENSIONS)
                if not is_supported and not is_image:
                    continue
                root_files_seen += 1

                if caption_only:
                    # Niente manifest, niente indice, niente EXIF/GPS: la cache
                    # (file + modello + prompt) decide cosa e' gia' stato fatto.
                    if not is_image or not index_images:
                        continue
                    try:
                        signature = get_file_signature(file_path)
                    except Exception as e:
                        logger.warning(f"Impossibile leggere gli attributi di {file_path}, saltato: {e}")
                        continue
                    if caption_cache.get(file_path, signature, vision_model, prompt_version) is not None:
                        cap_stats["cached"] += 1
                        continue
                    if is_likely_icon_or_graphic(file_path):
                        cap_stats["icons"] += 1
                        continue

                    t0 = time.time()
                    try:
                        caption = caption_image(
                            file_path, vision_model, prompt_version=prompt_version,
                            options=caption_options, max_side=caption_max_side,
                            timeout=caption_timeout,
                        )
                    except Exception as e:
                        logger.error(f"Errore captioning immagine {file_path}: {e}")
                        cap_stats["errors"] += 1
                        cap_stats["consecutive_errors"] += 1
                        if cap_stats["consecutive_errors"] >= MAX_CONSECUTIVE_ERRORS:
                            logger.error(
                                f"{MAX_CONSECUTIVE_ERRORS} errori consecutivi: interrompo. "
                                "Ollama e' in esecuzione e il modello vision e' scaricato?"
                            )
                            log_caption_only_summary()
                            return
                        continue

                    cap_stats["consecutive_errors"] = 0
                    if not caption:
                        logger.warning(f"  -> Didascalia vuota per {filename}, saltata.")
                        cap_stats["empty"] += 1
                        continue

                    caption_cache.put(file_path, signature, vision_model, prompt_version, caption)
                    cap_stats["generated"] += 1
                    cap_stats["gen_seconds"] += time.time() - t0
                    anteprima = caption[:120].replace("\n", " | ")
                    logger.info(f"  -> Didascalia generata: {filename}: {anteprima}...")
                    if cap_stats["generated"] % 25 == 0:
                        done = cap_stats["generated"] + cap_stats["cached"] + cap_stats["icons"] + cap_stats["empty"] + cap_stats["errors"]
                        avg = cap_stats["gen_seconds"] / cap_stats["generated"]
                        residue_h = max(total_images - done, 0) * avg / 3600
                        logger.info(
                            f"  == Progresso: {done}/{total_images} foto esaminate "
                            f"({cap_stats['generated']} generate, {cap_stats['cached']} in cache); "
                            f"{avg:.1f} s/foto; tempo residuo stimato al piu' {residue_h:.1f} h"
                        )
                    continue

                # Confronto con il manifest: se il file c'era gia' ed e'
                # rimasto identico (stessa data di modifica e dimensione),
                # lo saltiamo del tutto - i suoi chunk restano quelli gia'
                # presenti in Chroma dal run precedente.
                try:
                    signature = get_file_signature(file_path)
                except Exception as e:
                    logger.warning(f"Impossibile leggere gli attributi di {file_path}, saltato: {e}")
                    continue

                previous = old_manifest.get(file_path)
                # Una foto va ridescritta anche se il file e' identico quando la sua
                # didascalia e' stata generata con un altro modello vision o un'altra
                # versione di prompt (le voci vecchie, senza queste informazioni,
                # contano come diverse): altrimenti l'indice mescolerebbe didascalie
                # di stile e qualita' differenti.
                caption_stale = bool(
                    previous and is_image and index_images
                    and (previous.get("caption_model") != vision_model
                         or previous.get("caption_prompt") != prompt_version)
                )
                text_stale = bool(
                    previous and previous.get("text_version") != INDEX_TEXT_VERSION
                    and (is_supported or index_images)
                )
                if previous and not caption_stale and not text_stale and previous.get("mtime") == signature["mtime"] and previous.get("size") == signature["size"]:
                    new_manifest[file_path] = previous
                    unchanged_count += 1
                    continue

                if is_supported:
                    ext = os.path.splitext(filename)[1].lower()
                    full_text = extract_text_from_file(file_path, ext, logger)

                    if not full_text.strip():
                        new_manifest[file_path] = {**signature, "type": "empty", "text_version": INDEX_TEXT_VERSION}
                        checkpoint()
                        continue

                    doc = Document(text=full_text, metadata={"file_name": filename, "path": file_path, "type": "text"})
                    n_chunks = index_single_document(doc)
                    new_manifest[file_path] = {**signature, "type": "text", "text_version": INDEX_TEXT_VERSION}
                    checkpoint()
                    processed_count += 1
                    logger.info(f"  -> Indicizzato [{ext.upper()}] {filename} ({n_chunks} chunk)")

                elif is_image:
                    if not index_images:
                        continue
                    if is_likely_icon_or_graphic(file_path):
                        logger.info(f"  -> Immagine saltata (probabile icona/grafica, non foto): {filename}")
                        continue

                    # Didascalia dalla cache (stesso file, modello e prompt) oppure nuova.
                    caption = caption_cache.get(file_path, signature, vision_model, prompt_version)
                    caption_from_cache = caption is not None
                    if not caption_from_cache:
                        try:
                            caption = caption_image(
                                file_path, vision_model, prompt_version=prompt_version,
                                options=caption_options, max_side=caption_max_side,
                                timeout=caption_timeout,
                            )
                        except Exception as e:
                            logger.error(f"Errore captioning immagine {file_path}: {e}")
                            continue
                        if caption:
                            caption_cache.put(file_path, signature, vision_model, prompt_version, caption)

                    if not caption:
                        logger.warning(f"  -> Didascalia vuota per {filename}, saltata.")
                        continue

                    exif_info = extract_image_exif(file_path)
                    text_lines = []

                    if "gps_lat" in exif_info and "gps_lon" in exif_info:
                        match = find_nearest_location(exif_info["gps_lat"], exif_info["gps_lon"], gpx_points, max_distance_km=gpx_max_km)
                        if match and match["within_threshold"]:
                            exif_info["location_name"] = match["name"]
                            exif_info["region"] = match["region"]
                            exif_info["location_km"] = round(match["distance_km"], 1)
                            text_lines.append(f"Luogo: {match['name']} ({match['region']})")
                        else:
                            gap_entry = {
                                "file": filename,
                                "path": file_path,
                                "gps_lat": exif_info["gps_lat"],
                                "gps_lon": exif_info["gps_lon"],
                            }
                            if match:
                                gap_entry["nearest_known"] = match["name"]
                                gap_entry["nearest_region"] = match["region"]
                                gap_entry["nearest_distance_km"] = round(match["distance_km"], 1)
                            geocoding_gaps.append(gap_entry)

                    # Testo embeddato = solo cio' che ha un significato per la ricerca:
                    # luogo, nomi delle cartelle e descrizione. Data, coordinate,
                    # fotocamera e nome del file restano nei metadati (filtri e
                    # risultati), non nel testo.
                    folder = folder_label(root, target_path)
                    if folder:
                        text_lines.append(f"Cartella: {folder}")
                    text_lines.append(caption)
                    doc_text = "\n".join(text_lines)

                    metadata = {"file_name": filename, "path": file_path, "type": "image"}
                    metadata.update(exif_info)
                    if folder:
                        metadata["folder"] = folder
                    if "camera_model" in exif_info:
                        metadata["camera"] = camera_label(exif_info.get("camera_make"), exif_info["camera_model"])
                    date_info = parse_exif_date(exif_info.get("date_taken"))
                    if date_info:
                        metadata["year"], metadata["date_ts"] = date_info
                        metadata["year_source"] = "exif"
                    else:
                        metadata.pop("date_taken", None)  # data assente o non valida (es. 0000:00:00)
                        folder_year = year_from_folder(folder)
                        if folder_year:
                            metadata["year"] = folder_year
                            metadata["year_source"] = "cartella"
                    doc = Document(text=doc_text, metadata=metadata,
                                   excluded_embed_metadata_keys=list(metadata.keys()))
                    index_single_document(doc)
                    new_manifest[file_path] = {
                        **signature, "type": "image", "text_version": INDEX_TEXT_VERSION,
                        "caption_model": vision_model, "caption_prompt": prompt_version,
                    }
                    checkpoint()
                    processed_count += 1
                    origine = "da cache" if caption_from_cache else "generata"
                    anteprima = caption[:120].replace("\n", " | ")
                    logger.info(f"  -> Didascalia {origine} e indicizzata: {filename}: {anteprima}...")

            if not recursive:
                break

        (healthy_roots if root_files_seen else protected_roots).append(target_path)

    if caption_only:
        log_caption_only_summary()
        return

    # File presenti nel manifest precedente ma non ritrovati in questo giro:
    # sono stati spostati, rinominati o cancellati dal disco. Rimuoviamo i
    # loro vecchi chunk dall'indice per non lasciare risultati "fantasma".
    removed_paths = []
    kept_protected = 0
    for p in old_manifest:
        if p in new_manifest:
            continue
        if path_is_under(p, protected_roots) and not path_is_under(p, healthy_roots):
            new_manifest[p] = old_manifest[p]  # cartella non raggiungibile: si conserva tutto
            kept_protected += 1
        else:
            removed_paths.append(p)
    if kept_protected:
        logger.warning(
            f"  -> {kept_protected} file mantenuti nell'indice perche' la loro cartella radice non era "
            f"raggiungibile o non conteneva nessun file: {', '.join(protected_roots)}. Se l'hai svuotata "
            "di proposito, togli la cartella da 'directories' in config.json: al giro successivo i suoi "
            "file verranno rimossi dall'indice."
        )
    if removed_paths:
        logger.info(f"4. Rimozione dei chunk per {len(removed_paths)} file non piu' presenti sul disco...")
        for path in removed_paths:
            try:
                chroma_collection.delete(where={"path": path})
            except Exception as e:
                logger.warning(f"  -> Impossibile ripulire i chunk di {path}: {e}")

    save_manifest(manifest_path, new_manifest, logger)

    if geocoding_gaps:
        gaps_path = os.path.join(base_dir, "geocoding_gaps.json")
        try:
            with open(gaps_path, "w", encoding="utf-8") as f:
                json.dump(geocoding_gaps, f, ensure_ascii=False, indent=2)
            logger.info(
                f"  -> {len(geocoding_gaps)} foto con GPS senza localita' entro {gpx_max_km} km: "
                f"dettagli in '{gaps_path}'"
            )
        except Exception as e:
            logger.error(f"Impossibile salvare il report dei buchi di geocoding: {e}")

    if processed_count == 0 and not removed_paths:
        logger.info("Nessuna modifica rispetto all'ultima indicizzazione: nulla da aggiornare.")
    else:
        logger.info(
            f"Indicizzazione completata! {processed_count} file processati, {unchanged_count} invariati (saltati), "
            f"{len(removed_paths)} rimossi. Indice aggiornato in '{db_path}' ({len(new_manifest)} file totali tracciati)."
        )

if __name__ == "__main__":
    main()